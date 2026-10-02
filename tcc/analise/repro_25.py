from simlib import *
import openpyxl, glob, json
D='/tmp/claude-0/-home-user-Sistema-de-Suporte-a-Decisao/2289c1dd-334d-5daf-aca8-58cbfaaa581c/scratchpad/x/exp/exportacoes'
MM={m:i+1 for i,m in enumerate(MESES)}
out=[]
for f in sorted(glob.glob(D+'/*.xlsx')):
    wb=openpyxl.load_workbook(f,data_only=True); ws=wb['Parâmetros e notas']; p={}; cur=[]
    for r in range(1,ws.max_row+1):
        k=ws.cell(r,1).value
        if k in ('VM1','VM2','VM3'): cur.append((ws.cell(r,2).value,float(ws.cell(r,3).value),[float(ws.cell(r,c).value) for c in range(6,18)]))
        elif k: p[k]=ws.cell(r,2).value
    if p['Status da execução']!='Executado': continue
    per=p['Período'].replace(' a ','/').split('/')  # JAN/1911/DEZ/2021
    a=acude(p['Código']); cap=float(p['Capacidade no simulador (hm³)'])
    plano=[{"Faixa":e,"Racionamento":r,"NomeFaixaNormal":"Normal",**{m:v for m,v in zip(MESES,l)}} for e,r,l in cur]
    rr=dict(a, capacidade=cap, vol_inicial=cap*float(p['Volume inicial (%)'])/100, demanda=float(p['Q normal base (L/s)'])/1000, gatilho=0.0, plano_secas_custom=plano)
    d=list(simular([rr],ini=(per[0],int(per[1])),fim=(per[2],int(per[3])),niveis_meta=True).values())[0]
    ex=pd.read_excel(f,sheet_name='Resultados')
    dv=float(np.abs(ex['Armazenamento Final (hm³)'].values-d['Armazenamento Final'].values).max())
    de=int((ex['Modo Operação'].values!=d['Modo Operação'].values).sum())
    out.append((p['Hidrossistema'],len(d),dv,de,float(residuo(d).abs().max()))); print(out[-1])
json.dump(out,open('repro_25.json','w'))
print(max(o[2] for o in out), sum(o[3] for o in out), sum(o[1] for o in out))
